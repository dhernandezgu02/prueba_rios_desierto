from rest_framework import generics, status
from rest_framework.decorators import api_view
from rest_framework.response import Response
from django.shortcuts import get_object_or_404
from django.db.models import Q
from django.http import HttpResponse
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from datetime import date
from .models import TipoDocumento, Cliente, Compra
from .serializers import (
    TipoDocumentoSerializer, 
    ClienteSerializer, 
    ClientePerfilSerializer,  # Serializer para consulta de perfil completo
    CompraSerializer,
    CompraSimpleSerializer
)


@api_view(['GET'])
def consultar_cliente_por_documento(request, numero_documento):
    """
    API que consulta por número de documento la información básica del cliente.
    
    Devuelve los campos esenciales del perfil:
    - Numero de documento
    - Nombre
    - Apellido  
    - Correo
    - Teléfono
    
    La consulta se hace directamente en la base de datos SQLite.
    
    URL: /api/clientes/consulta/{numero_documento}/
    """
    try:
        # Buscar cliente activo por número de documento
        cliente = Cliente.objects.get(
            numero_documento=numero_documento.strip(),
            activo=True
        )
        
        # Usar el serializer específico para consulta de perfil
        serializer = ClientePerfilSerializer(cliente)
        
        return Response({
            'success': True,
            'data': serializer.data,
            'message': 'Cliente encontrado exitosamente'
        }, status=status.HTTP_200_OK)
        
    except Cliente.DoesNotExist:
        return Response({
            'success': False,
            'data': None,
            'message': f'No se encontró ningún cliente con número de documento: {numero_documento}'
        }, status=status.HTTP_404_NOT_FOUND)
    
    except Exception as e:
        return Response({
            'success': False,
            'data': None,
            'message': f'Error interno del servidor: {str(e)}'
        }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


@api_view(['GET'])
def buscar_cliente(request):
    """
    Busca un cliente por tipo de documento y número de documento.
    
    Parámetros requeridos:
    - tipo_documento: Código del tipo de documento (CC, NIT, CE, PP, TI)
    - numero_documento: Número de documento del cliente
    """
    tipo_documento_codigo = request.GET.get('tipo_documento')
    numero_documento = request.GET.get('numero_documento')
    
    if not tipo_documento_codigo or not numero_documento:
        return Response({
            'error': 'Se requieren los parámetros tipo_documento y numero_documento'
        }, status=status.HTTP_400_BAD_REQUEST)
    
    try:
        # Buscar tipo de documento
        tipo_documento = TipoDocumento.objects.get(
            codigo=tipo_documento_codigo.upper(),
            activo=True
        )
        
        # Buscar cliente
        cliente = Cliente.objects.get(
            tipo_documento=tipo_documento,
            numero_documento=numero_documento.strip(),
            activo=True
        )
        
        serializer = ClienteSerializer(cliente)
        return Response(serializer.data)
        
    except TipoDocumento.DoesNotExist:
        return Response({
            'error': f'Tipo de documento "{tipo_documento_codigo}" no válido'
        }, status=status.HTTP_400_BAD_REQUEST)
        
    except Cliente.DoesNotExist:
        return Response({
            'error': 'Cliente no encontrado'
        }, status=status.HTTP_404_NOT_FOUND)


@api_view(['GET'])
def compras_cliente(request, cliente_id):
    """
    Obtiene todas las compras de un cliente específico.
    """
    cliente = get_object_or_404(Cliente, id=cliente_id, activo=True)
    
    compras = Compra.objects.filter(
        cliente=cliente
    ).order_by('-fecha_compra')
    
    # Aplicar filtros opcionales
    estado = request.GET.get('estado')
    if estado:
        compras = compras.filter(estado=estado)
    
    # Paginación
    page_size = int(request.GET.get('page_size', 10))
    page = int(request.GET.get('page', 1))
    
    start = (page - 1) * page_size
    end = start + page_size
    
    compras_page = compras[start:end]
    total_count = compras.count()
    
    serializer = CompraSimpleSerializer(compras_page, many=True)
    
    return Response({
        'results': serializer.data,
        'count': total_count,
        'page': page,
        'page_size': page_size,
        'total_pages': (total_count + page_size - 1) // page_size
    })


class TipoDocumentoListView(generics.ListAPIView):
    """
    Lista todos los tipos de documento activos.
    """
    queryset = TipoDocumento.objects.filter(activo=True)
    serializer_class = TipoDocumentoSerializer


class ClienteListView(generics.ListAPIView):
    """
    Lista todos los clientes con filtros opcionales.
    """
    queryset = Cliente.objects.filter(activo=True)
    serializer_class = ClienteSerializer
    
    def get_queryset(self):
        queryset = super().get_queryset()
        
        # Filtros opcionales
        ciudad = self.request.GET.get('ciudad')
        if ciudad:
            queryset = queryset.filter(ciudad__icontains=ciudad)
        
        departamento = self.request.GET.get('departamento')
        if departamento:
            queryset = queryset.filter(departamento__icontains=departamento)
        
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(primer_nombre__icontains=search) |
                Q(primer_apellido__icontains=search) |
                Q(numero_documento__icontains=search) |
                Q(correo__icontains=search)
            )
        
        return queryset.order_by('-fecha_registro')


class ClienteDetailView(generics.RetrieveAPIView):
    """
    Detalle de un cliente específico.
    """
    queryset = Cliente.objects.filter(activo=True)
    serializer_class = ClienteSerializer


class CompraListView(generics.ListAPIView):
    """
    Lista todas las compras con filtros opcionales.
    """
    queryset = Compra.objects.all()
    serializer_class = CompraSerializer
    
    def get_queryset(self):
        queryset = super().get_queryset()
        
        # Filtros opcionales
        estado = self.request.GET.get('estado')
        if estado:
            queryset = queryset.filter(estado=estado)
        
        metodo_pago = self.request.GET.get('metodo_pago')
        if metodo_pago:
            queryset = queryset.filter(metodo_pago=metodo_pago)
        
        return queryset.order_by('-fecha_compra')


class CompraDetailView(generics.RetrieveAPIView):
    """
    Detalle de una compra específica.
    """
    queryset = Compra.objects.all()
    serializer_class = CompraSerializer


@api_view(['GET'])
def estadisticas_cliente(request, cliente_id):
    """
    Obtiene estadísticas de un cliente específico.
    """
    cliente = get_object_or_404(Cliente, id=cliente_id, activo=True)
    
    compras = cliente.compras.all()
    
    estadisticas = {
        'cliente': ClienteSerializer(cliente).data,
        'total_compras': compras.count(),
        'compras_completadas': compras.filter(estado='COMPLETADA').count(),
        'compras_pendientes': compras.filter(estado='PENDIENTE').count(),
        'monto_total': float(cliente.total_compras),
        'compra_mas_reciente': None,
        'compra_mas_antigua': None,
    }
    
    if compras.exists():
        compra_reciente = compras.order_by('-fecha_compra').first()
        compra_antigua = compras.order_by('fecha_compra').first()
        
        estadisticas['compra_mas_reciente'] = CompraSimpleSerializer(compra_reciente).data
        estadisticas['compra_mas_antigua'] = CompraSimpleSerializer(compra_antigua).data
    
    return Response(estadisticas)


@api_view(['GET'])
def reporte_fidelizacion_excel(request):
    """
    Genera un reporte Excel con clientes candidatos para fidelización.
    
    Criterios:
    - Clientes con compras del último mes >= $5,000,000 COP
    - Datos básicos del cliente + monto total último mes
    - Ordenados por monto de mayor a menor
    
    URL: /api/clientes/reporte/fidelizacion/
    """
    try:
        # Obtener parámetro opcional de monto mínimo
        monto_minimo_param = request.GET.get('monto_minimo', '5000000')
        try:
            monto_minimo = Decimal(monto_minimo_param)
        except (ValueError, TypeError):
            monto_minimo = Decimal('5000000')  # Default: 5 millones
        
        # Obtener candidatos para fidelización
        candidatos = Cliente.obtener_candidatos_fidelizacion(monto_minimo)
        
        if not candidatos:
            return Response({
                'success': False,
                'message': f'No se encontraron clientes con compras del último mes >= ${monto_minimo:,.0f} COP'
            }, status=status.HTTP_404_NOT_FOUND)
        
        # Crear archivo Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte Fidelización"
        
        # Configurar estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # Encabezados
        headers = [
            'ID Cliente',
            'Tipo Documento',
            'Número Documento', 
            'Nombre Completo',
            'Email',
            'Teléfono',
            'Ciudad',
            'Departamento',
            'Total Último Mes (COP)',
            'Cantidad Compras Último Mes',
            'Total Histórico (COP)',
            'Última Compra',
            'Estado Cliente'
        ]
        
        # Escribir encabezados
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # Escribir datos de candidatos
        for row, candidato in enumerate(candidatos, 2):
            cliente = candidato['cliente']
            
            ws.cell(row=row, column=1, value=cliente.id)
            ws.cell(row=row, column=2, value=cliente.tipo_documento.nombre)
            ws.cell(row=row, column=3, value=cliente.numero_documento)
            ws.cell(row=row, column=4, value=cliente.nombre_completo)
            ws.cell(row=row, column=5, value=cliente.correo)
            ws.cell(row=row, column=6, value=cliente.telefono)
            ws.cell(row=row, column=7, value=cliente.ciudad)
            ws.cell(row=row, column=8, value=cliente.departamento)
            ws.cell(row=row, column=9, value=float(candidato['total_ultimo_mes']))
            ws.cell(row=row, column=10, value=candidato['cantidad_compras_mes'])
            ws.cell(row=row, column=11, value=float(cliente.total_compras))
            ws.cell(row=row, column=12, value=cliente.ultima_compra.strftime('%d/%m/%Y') if cliente.ultima_compra else 'Sin compras')
            ws.cell(row=row, column=13, value='Activo' if cliente.activo else 'Inactivo')
        
        # Ajustar ancho de columnas
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
        
        # Agregar información del reporte
        ws['A' + str(len(candidatos) + 3)] = f"Reporte generado: {date.today().strftime('%d/%m/%Y')}"
        ws['A' + str(len(candidatos) + 4)] = f"Criterio mínimo: ${monto_minimo:,.0f} COP"
        ws['A' + str(len(candidatos) + 5)] = f"Total candidatos: {len(candidatos)}"
        
        # Preparar respuesta HTTP
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        timestamp = date.today().strftime('%Y%m%d')
        filename = f'reporte_fidelizacion_rios_desierto_{timestamp}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        # Guardar archivo en respuesta
        wb.save(response)
        
        return response
        
    except Exception as e:
        return Response({
            'success': False,
            'message': f'Error generando reporte: {str(e)}'
        }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
